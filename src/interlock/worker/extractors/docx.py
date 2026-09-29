"""Word document extractor - requires python-docx."""

from __future__ import annotations

import os

from interlock.worker.extractors.base import ExtractedContent

try:
    import docx as python_docx  # python-docx package

    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

_EXTENSIONS = {".docx"}


class DocxExtractor:
    """Extract text from Word .docx files using python-docx."""

    @property
    def available(self) -> bool:
        return _HAS_DOCX

    def can_handle(self, file_path: str, mime_type: str | None = None) -> bool:
        if not _HAS_DOCX:
            return False
        ext = os.path.splitext(file_path)[1].lower()
        return ext in _EXTENSIONS

    async def extract(self, file_path: str) -> ExtractedContent:
        if not _HAS_DOCX:
            return ExtractedContent(
                text="",
                metadata={"error": "python-docx not installed"},
            )

        doc = python_docx.Document(file_path)

        parts: list[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Extract table contents
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append("\t".join(cells))

        full_text = "\n".join(parts)
        return ExtractedContent(
            text=full_text,
            metadata={
                "source": file_path,
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
            },
            word_count=len(full_text.split()),
        )
